#!/usr/bin/env python3
"""Inspect values and formula caches read-only; never evaluate or write a workbook."""
import argparse
import json
import re
from pathlib import Path

FORMULA_TEXT = re.compile(r'^(?:[A-Z]{1,3}\d+\s*[+*/-]|(?:SUM|SUMPRODUCT|VLOOKUP|ROUND|IF|AVERAGE)\s*\()', re.I)


def looks_formula_text(value):
    return isinstance(value, str) and not value.lstrip().startswith('=') and bool(FORMULA_TEXT.match(value.strip()))


def label_cells(book):
    result = {}
    for sheet in book:
        for row in sheet:
            for cell in row:
                v = cell.value
                if isinstance(v, str) and cell.data_type != 'f' and not looks_formula_text(v) and cell.data_type != 'e':
                    result[(sheet.title, cell.coordinate)] = v.strip()
    return result


def inspect(path, baseline=None):
    try:
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError('需要openpyxl；可使用工作区提供的Python运行时，或已获授权的依赖环境') from exc
    book = load_workbook(path, data_only=False, read_only=True, keep_links=False)
    cache = load_workbook(path, data_only=True, read_only=True, keep_links=False)
    findings, notes = [], []
    expected = None
    if baseline:
        expected = load_workbook(baseline, data_only=False, read_only=True, keep_links=False)
        if book.sheetnames != expected.sheetnames:
            expected.close()
            expected = None
            notes.append('工作表列表不一致，拒绝自动模板比较；先核同版同布局')
        else:
            labels = label_cells(expected)
            current = label_cells(book)
            agreement = sum(current.get(k) == v for k, v in labels.items()) / len(labels) if labels else 0
            if agreement < 0.95:
                expected.close()
                expected = None
                notes.append('模板文本布局相似度不足95%，拒绝自动模板比较；相似也不证明公式正确')
    formula_count = 0
    try:
        for sheet in book:
            values = cache[sheet.title]
            baseline_sheet = expected[sheet.title] if expected else None
            # Iterate both sheets together: avoid repeatedly reopening streamed XML per cell.
            cached_rows = values.iter_rows()
            baseline_rows = baseline_sheet.iter_rows() if baseline_sheet else None
            for row_index, row in enumerate(sheet.iter_rows(), 1):
                cached_row = next(cached_rows, ())
                base_row = next(baseline_rows, ()) if baseline_rows else ()
                for i, cell in enumerate(row):
                    loc = {'sheet': sheet.title, 'cell': f'{get_column_letter(i + 1)}{row_index}'}
                    cached = cached_row[i].value if i < len(cached_row) else None
                    expected_cell = base_row[i] if i < len(base_row) else None
                    if cell.data_type == 'f':
                        formula_count += 1
                        if cached is None:
                            findings.append(dict(loc, kind='formula_cache_missing', formula=cell.value,
                                                 detail='没有可用缓存，不表示公式为0；openpyxl不重算'))
                        elif cached_row[i].data_type == 'e':
                            findings.append(dict(loc, kind='cached_error', formula=cell.value, value=cached))
                    elif looks_formula_text(cell.value):
                        findings.append(dict(loc, kind='formula_like_text', value=cell.value,
                                             detail='像缺少等号的公式文字；先核当前引用含义'))
                    elif cell.data_type == 'e':
                        findings.append(dict(loc, kind='literal_error', value=cell.value))
                    if expected_cell is not None and expected_cell.data_type == 'f' and cell.data_type != 'f':
                        findings.append(dict(loc, kind='formula_overridden', value=cell.value,
                                             template_formula=expected_cell.value,
                                             detail='同版模板公式被非公式值替代的候选；不自动恢复'))
                    elif expected_cell is not None and expected_cell.data_type == 'f' and cell.data_type == 'f' and cell.value != expected_cell.value:
                        findings.append(dict(loc, kind='formula_differs_from_template', formula=cell.value,
                                             template_formula=expected_cell.value,
                                             detail='与同版模板公式不同；可能是修正或引用错误，需要业务核实'))
    finally:
        book.close()
        cache.close()
        if expected:
            expected.close()
    return {'workbook_name': Path(path).name, 'read_only': True, 'formula_count': formula_count,
            'findings': findings, 'notes': notes + ['数字文字、缓存、模板比较仅是线索，需要业务与当前版本确认',
                                                  '原文件未修改；无数据时不填0，不重写最终权益']}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('workbook', type=Path)
    parser.add_argument('--baseline', type=Path)
    parser.add_argument('--out', type=Path, required=True)
    args = parser.parse_args()
    report = inspect(args.workbook, args.baseline)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(report, ensure_ascii=False, indent=2, default=str), encoding='utf-8')
    print(json.dumps({'read_only': True, 'formula_count': report['formula_count'], 'findings': len(report['findings'])}))


if __name__ == '__main__':
    main()
